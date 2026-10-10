"""Word 模板库：用 python-docx 生成 4 套文档模板。

模板里的 ``{{name}}`` 占位符由 ``zhixue.tools.docgen_tools.render_word`` 填充。
真实项目里把教研做好的 .docx 丢进 ``zhixue/data/templates`` 覆盖即可。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from zhixue.config import get_settings
from zhixue.logging_setup import get_logger

logger = get_logger("data.templates")

#: 模板清单：(template_id, 名称, 类型, 说明)
TEMPLATE_SPECS: list[tuple[str, str, str, str]] = [
    ("TPL-STUDY-PLAN", "学习计划模板", "study_plan", "按知识点排期的个性化学习计划"),
    ("TPL-LECTURE", "知识点讲义模板", "lecture", "含概念讲解、示例与课堂练习的讲义"),
    ("TPL-MISTAKE", "错题分析报告模板", "mistake_report", "围绕错题与薄弱知识点的分析报告"),
    ("TPL-EXAM", "试卷模板", "exam_paper", "按题型分组的试卷排版"),
]


def _build_study_plan(doc) -> None:
    doc.add_heading("{{title}}", level=0)
    doc.add_paragraph("学生：{{student_name}}　　班级：{{class_name}}　　生成日期：{{date}}")
    doc.add_heading("一、学习目标", level=1)
    doc.add_paragraph("{{goal}}")
    doc.add_heading("二、知识点安排", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "序号"
    header[1].text = "知识点"
    header[2].text = "建议用时"
    doc.add_paragraph("{{schedule_rows}}")
    doc.add_heading("三、教师寄语", level=1)
    doc.add_paragraph("{{remark}}")


def _build_lecture(doc) -> None:
    doc.add_heading("{{title}}", level=0)
    doc.add_paragraph("学科：{{subject}}　知识点：{{knowledge_point}}　难度：{{difficulty}}")
    doc.add_heading("一、概念讲解", level=1)
    doc.add_paragraph("{{explanation}}")
    doc.add_heading("二、典型示例", level=1)
    doc.add_paragraph("{{example}}")
    doc.add_heading("三、课堂练习", level=1)
    doc.add_paragraph("{{exercises}}")
    doc.add_heading("四、易错提醒", level=1)
    doc.add_paragraph("{{pitfall}}")


def _build_mistake_report(doc) -> None:
    doc.add_heading("{{title}}", level=0)
    doc.add_paragraph("学生：{{student_name}}　统计范围：{{range_desc}}　生成日期：{{date}}")
    doc.add_heading("一、整体表现", level=1)
    doc.add_paragraph("{{summary}}")
    doc.add_heading("二、错题明细", level=1)
    doc.add_paragraph("{{mistake_rows}}")
    doc.add_heading("三、薄弱知识点", level=1)
    doc.add_paragraph("{{weak_points}}")
    doc.add_heading("四、提升建议", level=1)
    doc.add_paragraph("{{advice}}")


def _build_exam(doc) -> None:
    doc.add_heading("{{title}}", level=0)
    doc.add_paragraph("学科：{{subject}}　总分：{{full_score}}　建议用时：{{duration}} 分钟")
    doc.add_paragraph("姓名：__________　班级：__________　得分：__________")
    doc.add_heading("一、试题", level=1)
    doc.add_paragraph("{{questions}}")
    doc.add_heading("二、参考答案", level=1)
    doc.add_paragraph("{{answers}}")


_BUILDERS = {
    "study_plan": _build_study_plan,
    "lecture": _build_lecture,
    "mistake_report": _build_mistake_report,
    "exam_paper": _build_exam,
}


def build_templates(template_dir: Path | None = None, force: bool = False) -> list[dict[str, Any]]:
    """生成模板文件，返回模板元数据列表（可直接写库）。"""
    from docx import Document

    settings = get_settings()
    template_dir = template_dir or settings.template_dir
    template_dir.mkdir(parents=True, exist_ok=True)

    records: list[dict[str, Any]] = []
    for template_id, name, doc_type, description in TEMPLATE_SPECS:
        path = template_dir / f"{doc_type}.docx"
        if force or not path.exists():
            doc = Document()
            _BUILDERS[doc_type](doc)
            doc.save(path)
            logger.info("生成 Word 模板: %s", path.name)
        records.append(
            {
                "template_id": template_id,
                "name": name,
                "doc_type": doc_type,
                "path": str(path),
                "description": description,
            }
        )
    return records


__all__ = ["TEMPLATE_SPECS", "build_templates"]
