"""文档生成 Agent 的工具集（对应架构图 MCP DocGen :8010/mcp）。

* ``search_textbook``  -> RAG 课本检索（Milvus 向量库 / 关键词降级）
* ``render_word``      -> Word 模板生成（python-docx 填充 Word 模板库）
"""

from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path
from typing import Any

from zhixue.config import get_settings
from zhixue.logging_setup import get_logger
from zhixue.store import get_mysql
from zhixue.vector import get_retriever

logger = get_logger("tools.docgen")


# --------------------------------------------------------------------------- #
# 课本检索
# --------------------------------------------------------------------------- #
def search_textbook(
    query: str, subject: str | None = None, k: int | None = None
) -> list[dict[str, Any]]:
    """在课本向量库里检索与 ``query`` 最相关的段落。"""
    settings = get_settings()
    k = k or settings.retrieval_k
    hits = get_retriever().search_textbook(query, k=k, subject=subject or None)
    logger.info("课本检索 '%s' -> %d 条 (subject=%s)", query, len(hits), subject)
    return [
        {
            "chunk_id": h.get("id"),
            "title": h.get("title", ""),
            "chapter": h.get("chapter", ""),
            "subject": h.get("subject", ""),
            "kp_id": h.get("kp_id", ""),
            "content": h.get("_row", {}).get("content") or h.get("text", ""),
            "score": round(float(h.get("score", 0.0)), 4),
            "backend": h.get("backend", ""),
        }
        for h in hits
    ]


# --------------------------------------------------------------------------- #
# Word 模板
# --------------------------------------------------------------------------- #
def list_templates() -> list[dict[str, Any]]:
    """列出 Word 模板库。"""
    rows = get_mysql().query(
        "SELECT template_id, name, doc_type, path, description FROM word_template ORDER BY template_id"
    )
    return [dict(r) for r in rows]


def _find_template(doc_type: str) -> dict[str, Any]:
    row = get_mysql().query_one(
        "SELECT template_id, name, doc_type, path FROM word_template WHERE doc_type = %s",
        (doc_type,),
    )
    if row is None:
        available = [t["doc_type"] for t in list_templates()]
        raise ValueError(f"未找到类型为 {doc_type!r} 的模板，可用类型: {available}")
    return dict(row)


def _replace_paragraph(paragraph, mapping: dict[str, str]) -> None:
    """在保留原有格式的前提下替换 ``{{key}}``。"""
    if not paragraph.runs:
        return
    text = "".join(run.text for run in paragraph.runs)
    if "{{" not in text:
        return
    for key, value in mapping.items():
        text = text.replace("{{" + key + "}}", str(value))
    # 把结果写回第一个 run，其余清空，避免破坏字体设置
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _fill_document(document, mapping: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        _replace_paragraph(paragraph, mapping)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_paragraph(paragraph, mapping)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                _replace_paragraph(paragraph, mapping)


def render_word(
    doc_type: str,
    data: dict[str, Any],
    output_name: str | None = None,
) -> dict[str, Any]:
    """用 ``data`` 填充模板，生成一份 .docx。

    :param doc_type: ``study_plan`` / ``lecture`` / ``mistake_report`` / ``exam_paper``
    :param data: 占位符 -> 内容
    :param output_name: 输出文件名（不含目录），默认按类型加时间戳
    """
    from docx import Document

    settings = get_settings()
    template = _find_template(doc_type)
    template_path = Path(template["path"])
    if not template_path.exists():
        raise FileNotFoundError(f"模板文件不存在: {template_path}")

    mapping = {k: v for k, v in data.items()}
    mapping.setdefault("date", datetime.now().strftime("%Y-%m-%d"))
    mapping.setdefault("title", template["name"])

    document = Document(str(template_path))
    _fill_document(document, mapping)

    settings.output_dir.mkdir(parents=True, exist_ok=True)
    filename = output_name or f"{doc_type}_{int(time.time())}.docx"
    if not filename.endswith(".docx"):
        filename += ".docx"
    output_path = settings.output_dir / filename
    document.save(str(output_path))
    logger.info("生成 Word 文档: %s", output_path)
    return {
        "doc_type": doc_type,
        "template": template["name"],
        "path": str(output_path),
        "filename": output_path.name,
        "size": output_path.stat().st_size,
    }


# --------------------------------------------------------------------------- #
# 组合用法
# --------------------------------------------------------------------------- #
def build_study_plan(
    topic: str,
    subject: str | None = None,
    student_name: str = "同学",
    class_name: str = "—",
    days: int = 7,
    goal: str | None = None,
) -> dict[str, Any]:
    """检索课本 -> 组装学习计划 -> 渲染 Word。"""
    chunks = search_textbook(topic, subject=subject, k=4)
    schedule = []
    for index, chunk in enumerate(chunks, start=1):
        schedule.append(
            f"第 {index} 天：{chunk['title']}｜建议 {max(30, 90 - index * 10)} 分钟｜"
            f"来源：{chunk['chapter'] or chunk['subject']}"
        )
    if not schedule:
        schedule = [f"第 {index} 天：{topic} 强化练习｜建议 60 分钟" for index in range(1, 4)]

    return render_word(
        "study_plan",
        {
            "title": f"{topic} · 学习计划",
            "student_name": student_name,
            "class_name": class_name,
            "goal": goal or f"{days} 天内掌握「{topic}」的核心概念并能独立完成配套练习。",
            "schedule_rows": "\n".join(schedule),
            "remark": "按计划执行，每完成一天就在错题本上记录一次自测结果。",
        },
        output_name=f"study_plan_{int(time.time())}.docx",
    )


def build_lecture_note(
    knowledge_point: str,
    subject: str | None = None,
    explanation: str = "",
    example: str = "",
    exercises: str = "",
    pitfall: str = "",
    difficulty: str = "中等",
) -> dict[str, Any]:
    """生成知识点讲义。"""
    return render_word(
        "lecture",
        {
            "title": f"{knowledge_point} 讲义",
            "subject": subject or "—",
            "knowledge_point": knowledge_point,
            "difficulty": difficulty,
            "explanation": explanation,
            "example": example,
            "exercises": exercises,
            "pitfall": pitfall,
        },
        output_name=f"lecture_{int(time.time())}.docx",
    )


__all__ = [
    "search_textbook",
    "list_templates",
    "render_word",
    "build_study_plan",
    "build_lecture_note",
]
